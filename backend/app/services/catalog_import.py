"""Katalogni Excel (.xlsx) yoki CSV dan import qilish.

Har bir qator — bitta variant. Bir xil nomli qatorlar bitta mahsulotga birlashadi.
Razmer katakchasida "S, M, L" yozilsa, har biri alohida variant bo'ladi.
"""

import csv
import io
import re
from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models import Product
from app.services.catalog import save_product
from app.services.text import normalize

TEMPLATE_HEADERS = ["kategoriya", "nomi", "narx", "tavsif", "razmer", "rang", "qoldiq", "rasm"]

HEADER_ALIASES = {
    "category": ("kategoriya", "category", "kategoriya nomi", "kategoriya", "категория"),
    "name": ("nomi", "nom", "name", "mahsulot", "название", "наименование", "товар"),
    "price": ("narx", "narxi", "price", "цена", "нарх"),
    "description": ("tavsif", "description", "описание", "izoh"),
    "size": ("razmer", "olcham", "size", "размер", "ўлчам"),
    "color": ("rang", "color", "цвет", "ранг"),
    "stock": ("qoldiq", "soni", "stock", "остаток", "количество", "қолдиқ"),
    "images": ("rasm", "rasmlar", "image", "images", "photo", "фото", "расм"),
    "sku": ("sku", "artikul", "артикул"),
}
_ALIAS_LOOKUP = {normalize(alias): key for key, aliases in HEADER_ALIASES.items() for alias in aliases}


@dataclass
class ImportResult:
    created: int = 0
    updated: int = 0
    errors: list[str] = field(default_factory=list)
    product_ids: list[int] = field(default_factory=list)


def _read_rows(content: bytes, filename: str) -> list[list]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        return [list(r) for r in wb.active.iter_rows(values_only=True)]
    text = content.decode("utf-8-sig", errors="replace")
    dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t") if text.strip() else csv.excel
    return [list(r) for r in csv.reader(io.StringIO(text), dialect)]


def _to_int(value) -> int | None:
    """250000, "250 000", "250,000", "250.000 so'm" → 250000."""
    if value is None:
        return None
    if isinstance(value, int | float):
        return int(value)
    text = str(value).strip().replace(" ", "").replace("\u00a0", "")
    text = re.sub(r"(so'?m|сум|uzs)$", "", text, flags=re.IGNORECASE)
    if re.fullmatch(r"\d{1,3}([.,]\d{3})+", text):
        return int(re.sub(r"[.,]", "", text))
    m = re.match(r"\d+", text)
    return int(m.group()) if m else None


def _split(value) -> list[str]:
    raw = str(value or "").replace(";", ",").replace("\n", ",")
    return [x.strip() for x in raw.split(",") if x.strip()]


def parse_rows(rows: list[list]) -> tuple[list[dict], list[str]]:
    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        return [], ["Fayl bo'sh"]
    header = [_ALIAS_LOOKUP.get(normalize(str(h or ""))) for h in rows[0]]
    if "name" not in header or "price" not in header:
        return [], ["Sarlavhada kamida 'nomi' va 'narx' ustunlari bo'lishi kerak"]

    products: dict[str, dict] = {}
    errors: list[str] = []
    for line_no, row in enumerate(rows[1:], start=2):
        rec = {key: row[i] if i < len(row) else None for i, key in enumerate(header) if key}
        name = str(rec.get("name") or "").strip()
        price = _to_int(rec.get("price"))
        if not name:
            errors.append(f"{line_no}-qator: nomi yo'q")
            continue
        key = normalize(name)
        prod = products.get(key)
        if prod is None:
            if price is None:
                errors.append(f"{line_no}-qator: narx noto'g'ri")
                continue
            prod = products[key] = {
                "name": name,
                "price": price,
                "description": str(rec.get("description") or "").strip(),
                "category": str(rec.get("category") or "").strip() or None,
                "variants": [],
                "images": [],
            }
        for url in _split(rec.get("images")):
            if url.startswith(("http://", "https://")) and url not in prod["images"]:
                prod["images"].append(url)
        sizes = _split(rec.get("size")) or [None]
        color = str(rec.get("color") or "").strip() or None
        stock = _to_int(rec.get("stock"))
        for size in sizes:
            prod["variants"].append(
                {
                    "attrs": {"size": size, "color": color},
                    "sku": str(rec.get("sku") or "").strip() or None,
                    "price_override": price if price is not None and price != prod["price"] else None,
                    "stock": stock,
                }
            )
    return list(products.values()), errors


async def import_catalog(session: AsyncSession, shop_id: int, content: bytes, filename: str) -> ImportResult:
    result = ImportResult()
    try:
        rows = _read_rows(content, filename)
    except Exception as exc:  # noqa: BLE001 — foydalanuvchiga tushunarli xato qaytaramiz
        result.errors.append(f"Faylni o'qib bo'lmadi: {exc}")
        return result
    parsed, result.errors = parse_rows(rows)

    existing = {
        normalize(p.name): p
        for p in (
            await session.scalars(
                select(Product)
                .where(Product.shop_id == shop_id)
                .options(selectinload(Product.variants), selectinload(Product.images))
            )
        ).all()
    }
    for item in parsed:
        product = existing.get(normalize(item["name"]))
        saved = await save_product(
            session,
            shop_id,
            product=product,
            name=item["name"],
            price=item["price"],
            description=item["description"],
            category_name=item["category"],
            variants=item["variants"],
            images=item["images"] or (None if product else []),
        )
        result.product_ids.append(saved.id)
        if product:
            result.updated += 1
        else:
            result.created += 1
    return result


def template_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "katalog"
    ws.append(TEMPLATE_HEADERS)
    ws.append(["Ko'ylaklar", "Qora ko'ylak", 250000, "Paxta, yozgi", "42, 44", "qora", 3, "https://..."])
    ws.append(["Ko'ylaklar", "Qora ko'ylak", 250000, "", "46", "qora", 1, ""])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
